import glob
import pptx
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

# Search for the PPTX file in current directory
candidates = glob.glob("*.pptx")
target_file = None
for f in candidates:
  if "NovaBank" in f or "Appendix" in f:
    target_file = f
    break
if not target_file and candidates:
  target_file = candidates[0]

if not target_file:
  print("Error: No .pptx file found in this folder.")
  exit()

print(f"Modifying presentation: {target_file}...")
prs = Presentation(target_file)

# -------------------------------------------------------------
# 1. Slide 1: Ensure GitHub repository link is present
# -------------------------------------------------------------
slide1 = prs.slides[0]
for shape in slide1.shapes:
  if shape.has_text_frame and "A Data-Driven Strategy" in shape.text_frame.text:
    if "github.com" not in shape.text_frame.text:
      p = shape.text_frame.add_paragraph()
      p.text = (
          "Reproducible Repository:"
          " https://github.com/popicu/novabank-retention-framework"
      )
      p.font.size = Pt(11)
      p.font.color.rgb = RGBColor(43, 108, 176)
      p.font.bold = True
      print("Slide 1: Added GitHub link.")

# -------------------------------------------------------------
# 2. Slide 10: Complete Redesign for AI Usage & Governance
# -------------------------------------------------------------
slide10 = prs.slides[9]

# Clean up any old placeholder text frames on Slide 10
shapes_to_delete = []
for idx, s in enumerate(slide10.shapes):
  if idx > 0 and s.has_text_frame:  # Keep background picture/title
    if "Ready for Executive" in s.text_frame.text or "Team |" in s.text_frame.text:
      shapes_to_delete.append(s)

for s in shapes_to_delete:
  sp_elm = s._element
  sp_elm.getparent().remove(sp_elm)

# Subtitle
sub_box = slide10.shapes.add_textbox(
    Inches(0.6), Inches(1.5), Inches(12.0), Inches(0.4)
)
p_sub = sub_box.text_frame.paragraphs[0]
p_sub.text = (
    "Audit Defensibility, Human-in-the-Loop AI Attribution, and Continuous"
    " Operational Oversight"
)
p_sub.font.size = Pt(13)
p_sub.font.color.rgb = RGBColor(74, 85, 104)

# Panel 1 (Left): AI Copilot Attribution & Governance
p1 = slide10.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(0.6),
    Inches(2.1),
    Inches(5.8),
    Inches(3.7),
)
p1.fill.solid()
p1.fill.fore_color.rgb = RGBColor(248, 250, 252)
p1.line.color.rgb = RGBColor(203, 213, 224)
tf1 = p1.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_right = Inches(0.25)
tf1.margin_top = tf1.margin_bottom = Inches(0.2)

p_h1 = tf1.paragraphs[0]
p_h1.text = "🤖 AI Copilot Attribution & Governance"
p_h1.font.bold = True
p_h1.font.size = Pt(12)
p_h1.font.color.rgb = RGBColor(15, 41, 66)

bullets1 = [
    (
        "Assistance Scope",
        (
            "Generative AI copilots supported scikit-learn syntax boilerplate,"
            " Matplotlib chart formatting, and Markdown scaffolding."
        ),
    ),
    (
        "Human Direction",
        (
            "Core decision logic, target engineering, business objective"
            " function (€15/call vs €250 CLV), and decile policy rules were"
            " designed by the project author."
        ),
    ),
    (
        "Data Leakage Remediation",
        (
            "Independently identified and audited the 'duration' feature trap,"
            " mandating its drop to preserve pre-call operational validity."
        ),
    ),
]
for title, desc in bullets1:
  p = tf1.add_paragraph()
  p.text = f"• {title}: {desc}"
  p.font.size = Pt(9.5)
  p.font.color.rgb = RGBColor(45, 55, 72)
  p.space_before = Pt(6)

# Panel 2 (Right): Model Governance & Operational Controls
p2 = slide10.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(6.8),
    Inches(2.1),
    Inches(5.8),
    Inches(3.7),
)
p2.fill.solid()
p2.fill.fore_color.rgb = RGBColor(248, 250, 252)
p2.line.color.rgb = RGBColor(203, 213, 224)
tf2 = p2.text_frame
tf2.word_wrap = True
tf2.margin_left = tf2.margin_right = Inches(0.25)
tf2.margin_top = tf2.margin_bottom = Inches(0.2)

p_h2 = tf2.paragraphs[0]
p_h2.text = "🛡️ Model Governance & Operational Controls"
p_h2.font.bold = True
p_h2.font.size = Pt(12)
p_h2.font.color.rgb = RGBColor(15, 41, 66)

bullets2 = [
    (
        "Macro Drift Monitoring",
        (
            "Because macro variables (Euribor 3M, Emp.Var) govern 68% of"
            " propensity variance, models undergo mandatory quarterly"
            " recalibration."
        ),
    ),
    (
        "Fairness & Demographic Neutrality",
        (
            "Tree SHAP confirms zero decision exploitation of protected"
            " attributes (age, marital, education), ensuring regulatory"
            " compliance."
        ),
    ),
    (
        "Contact Saturation Ceiling",
        (
            "Strict guardrail capping outreach at 3 outbound calls per customer"
            " per 90 days to eliminate brand fatigue and churn risk."
        ),
    ),
]
for title, desc in bullets2:
  p = tf2.add_paragraph()
  p.text = f"• {title}: {desc}"
  p.font.size = Pt(9.5)
  p.font.color.rgb = RGBColor(45, 55, 72)
  p.space_before = Pt(6)

# Bottom Verification Banner
banner = slide10.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(0.6),
    Inches(6.0),
    Inches(12.0),
    Inches(0.85),
)
banner.fill.solid()
banner.fill.fore_color.rgb = RGBColor(235, 248, 255)
banner.line.color.rgb = RGBColor(43, 108, 176)
banner.line.width = Pt(1.5)
tf_b = banner.text_frame
tf_b.word_wrap = True
tf_b.margin_left = Inches(0.25)
tf_b.margin_top = Inches(0.12)

p_b1 = tf_b.paragraphs[0]
p_b1.text = "🔗 Reproducible Codebase & Auditable Submission Artifacts"
p_b1.font.bold = True
p_b1.font.size = Pt(10.5)
p_b1.font.color.rgb = RGBColor(15, 41, 66)

p_b2 = tf_b.add_paragraph()
p_b2.text = (
    "Public GitHub Repository:"
    " https://github.com/popicu/novabank-retention-framework  |  Includes"
    " main.py, reproducible notebook, requirements.txt, and diagnostic charts."
)
p_b2.font.size = Pt(9)
p_b2.font.color.rgb = RGBColor(43, 108, 176)
p_b2.font.bold = True
print("Slide 10: Governance panels and GitHub verification banner installed.")

# -------------------------------------------------------------
# 3. Slide 11: Remove redundant single-link image source slide
# -------------------------------------------------------------
if len(prs.slides) > 10:
  rId = prs.slides._sldIdLst[10].rId
  prs.part.drop_rel(rId)
  del prs.slides._sldIdLst[10]
  print("Slide 11: Redundant image slide deleted.")

output_name = "NovaBank_Predictive_Retention_Deck_Final.pptx"
prs.save(output_name)
print(
    f"SUCCESS: Generated '{output_name}' with exactly {len(prs.slides)} slides!"
)